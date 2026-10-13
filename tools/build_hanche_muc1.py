"""Redo the illustration of limitation 01 in 'Hạn chế.pptx' (slide 3: current, slide 4: optimized).

Keeps the deck's own map (vertical racks A..T, 2 cross aisles, inbound top, outbound bottom) and replaces
the cartons / numbers / routes with one TO of 9 bins whose two orders follow the routing app:
- current ("sweepOrder"): aisles in A→T order, only aisles with bins, each aisle swept one way from the end
  nearer the truck; aisle changes use the nearest crossing (front, cross aisle or far end).
- optimized: same bins re-sequenced (1-3-4-7-8-9-6-5-2), nearby bins handled back to back.
Each step is a group that fades in on click.
Usage: python3 tools/build_hanche_muc1.py <in.pptx> <out.pptx>
"""
import io
import sys
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SRC, DST = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

# ---- map geometry (EMU), read from the deck's own shapes ----
COL0, PITCH = 6667500, 482600              # rack A centre, rack-to-rack pitch
AISLE = lambda k: COL0 + (k + 0.5) * PITCH  # aisle k sits between rack column k and k+1
Y_TOP = 2146300                            # far end (inbound side), between rack tags and racks
Y_C1, Y_C2 = 3213100, 4406900              # cross aisle centres
Y_FRONT = 5511800                          # front aisle, between racks and staging
DOOR_X = 6400000                           # forklift at the staging, left
Y_DOOR = 5640000

# 9 bins: (aisle index, y). Aisles: 0 A|B, 2 C|D, 3 D|…, 6 Q|R, 7 R|S.
BIN = {
    1: (0, 3810000),
    2: (2, 5110000),
    3: (2, 2420000),
    4: (3, 2560000),
    5: (3, 5110000),
    6: (6, 5110000),
    7: (6, 2460000),
    8: (7, 3810000),
    9: (7, 5110000),
}
P = {n: (AISLE(k), y) for n, (k, y) in BIN.items()}
CROSS = [Y_TOP, Y_C1, Y_C2, Y_FRONT]


def leg(a, b):
    """Rectilinear path between two points in aisles, crossing at the nearest crossing (same rule as the app)."""
    (x1, y1), (x2, y2) = a, b
    if abs(x1 - x2) < 1:
        return [a, b]
    yc = min(CROSS, key=lambda c: abs(y1 - c) + abs(y2 - c))
    return [a, (x1, yc), (x2, yc), b]


def door_out(b):
    x, y = b
    return [(DOOR_X, Y_DOOR), (DOOR_X, Y_FRONT), (x, Y_FRONT), (x, y)] if y != Y_FRONT else [(DOOR_X, Y_DOOR), (DOOR_X, Y_FRONT), (x, y)]


def door_back(a, dy=0):
    x, y = a
    return [(x, y), (x, Y_FRONT + dy), (DOOR_X + 60000, Y_FRONT + dy), (DOOR_X + 60000, Y_DOOR)]


def length(order):
    pts = door_out(P[order[0]])
    for a, b in zip(order, order[1:]):
        pts += leg(P[a], P[b])[1:]
    pts += door_back(P[order[-1]])[1:]
    return sum(abs(p[0] - q[0]) + abs(p[1] - q[1]) for p, q in zip(pts, pts[1:]))


CUR = [1, 2, 3, 4, 5, 6, 7, 8, 9]
OPT = [1, 3, 4, 7, 8, 9, 6, 5, 2]
L_CUR, L_OPT = length(CUR), length(OPT)
assert L_OPT < L_CUR, (L_CUR, L_OPT)
SAVE = round(100 * (L_CUR - L_OPT) / L_CUR)

NAVY, RED, GREEN, INK = '0B2470', 'D7263D', '03A176', '1F2A44'
BLUE = '1F5FBF'


def rgb(h):
    return RGBColor.from_string(h)


def no_shadow(sh):
    sp = sh._element.spPr
    if sp.find(qn('a:effectLst')) is None:
        etree.SubElement(sp, qn('a:effectLst'))


def shape(sc, kind, x, y, w, h, fill=None, line=None, lw=1.0, dash=None, adj=None):
    s = sc.add_shape(kind, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    if adj is not None:
        s.adjustments[0] = adj
    no_shadow(s)
    return s


def text_into(s, lines, align=PP_ALIGN.CENTER, inset=30000):
    tf = s.text_frame
    tf.word_wrap = True
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Emu(inset))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, parts in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, sz, c, b in parts:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = rgb(c); r.font.name = 'Arial'


def label(sc, x, y, w, h, lines, color, fill='FFFFFF'):
    s = shape(sc, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, color, 1.25, adj=0.25)
    text_into(s, lines)
    return s


def poly(sc, pts, color, w=2.25, arrow=True, dash=None):
    fb = sc.build_freeform(int(pts[0][0]), int(pts[0][1]), scale=1.0)
    fb.add_line_segments([(int(x), int(y)) for x, y in pts[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = rgb(color); s.line.width = Pt(w)
    if dash:
        s.line.dash_style = dash
    no_shadow(s)
    ln = s._element.spPr.find(qn('a:ln'))
    etree.SubElement(ln, qn('a:round'))
    if arrow:
        t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return s


def route(order_pairs, start_door=False, end_door=False, back_dy=0):
    pts = []
    for a, b in order_pairs:
        seg = (door_out(P[b]) if a == 'door' else leg(P[a], P[b]))
        pts += seg if not pts else seg[1:]
    if end_door:
        seg = door_back(P[order_pairs[-1][1]] if order_pairs else None, back_dy)
        pts += seg if not pts else seg[1:]
    return pts


def chain(seq, first_from_door=False):
    pairs = list(zip(seq, seq[1:]))
    return ([('door', seq[0])] if first_from_door else []) + pairs


CARTON = None
for sh in prs.slides[2].shapes:
    if sh.name.startswith('P1 carton'):
        CARTON = sh.image.blob
        break
assert CARTON


def bins(slide, name, color, numbers):
    g = slide.shapes.add_group_shape(); g.name = name
    for n, (x, y) in P.items():
        g.shapes.add_picture(io.BytesIO(CARTON), Emu(int(x - 200000)), Emu(int(y - 75000)), Emu(165100), Emu(149622))
        b = shape(g.shapes, MSO_SHAPE.OVAL, x + 35000, y - 175000, 170000, 170000, color, 'FFFFFF', 1.0)
        text_into(b, [[(str(numbers[n]), 7.5, 'FFFFFF', True)]], inset=0)
    return g


def ring(sc, a, b, color):
    (x1, y1), (x2, y2) = P[a], P[b]
    x0, x3 = min(x1, x2) - 260000, max(x1, x2) + 260000
    return shape(sc, MSO_SHAPE.OVAL, x0, y1 - 250000, x3 - x0, 400000, None, color, 1.75, dash=MSO_LINE.DASH)


def clear(slide):
    for sh in list(slide.shapes):
        if sh.name.startswith(('P1 carton', 'PB ', 'Route 1', 'Route OPT')) and sh.name != '!!Route 1 forklift 1':
            sh._element.getparent().remove(sh._element)
    t = slide._element.find(qn('p:timing'))
    if t is not None:
        slide._element.remove(t)


PNS = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def set_timing(slide, steps):
    ids = iter(range(3, 100000))
    pars = []
    for step in steps:
        inner = []
        for k, sh in enumerate(step):
            a, b, c = next(ids), next(ids), next(ids)
            node = 'clickEffect' if k == 0 else 'withEffect'
            inner.append(
                f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="500"/><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
        o1, o2 = next(ids), next(ids)
        pars.append(f'<p:par><p:cTn id="{o1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                    f'<p:par><p:cTn id="{o2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>' + ''.join(inner) +
                    '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{PNS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + ''.join(pars) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml)
    ext = slide._element.find(qn('p:extLst'))
    if ext is not None:
        ext.addprevious(el)
    else:
        slide._element.append(el)


def group(slide, name):
    g = slide.shapes.add_group_shape(); g.name = name
    return g


# Callout boxes sit in the cross-aisle bands (empty horizontal strips).
BAND1 = (Y_C1 - 125000, 250000)
BAND2 = (Y_C2 - 125000, 250000)

# ================= Slide 3: current =================
s = prs.slides[2]
clear(s)
g1 = bins(s, 'Bước 1 – 9 thùng theo phiếu', BLUE, {n: n for n in P})
note = label(g1.shapes, 6280000, BAND1[0] - 40000, 1530000, BAND1[1] + 80000, [[('Phiếu in theo dãy A → T', 7, INK, True)], [('#1 → #2 → … → #9', 7, BLUE, True)]], BLUE)
g2 = group(s, 'Bước 2 – Cửa → #1 → #2')
poly(g2.shapes, route(chain([1, 2], True)), RED)
g3 = group(s, 'Bước 3 – #2 → #3 → #4 → #5')
poly(g3.shapes, route(chain([2, 3, 4, 5])), RED)
ring(g3.shapes, 2, 5, RED)
label(g3.shapes, 8420000, BAND2[0] - 40000, 1330000, BAND2[1] + 80000, [[('#2 · #5 sát nhau', 7, RED, True)], [('không hạ liên tiếp', 6.5, INK, False)]], RED)
g4 = group(s, 'Bước 4 – #5 → … → #9')
poly(g4.shapes, route(chain([5, 6, 7, 8, 9])), RED)
ring(g4.shapes, 6, 9, RED)
label(g4.shapes, 10340000, BAND2[0] - 40000, 1110000, BAND2[1] + 80000, [[('#6 · #9 sát nhau', 7, RED, True)], [('không hạ liên tiếp', 6.5, INK, False)]], RED)
g5 = group(s, 'Bước 5 – Về cửa + kết luận')
poly(g5.shapes, door_back(P[9], 70000), RED)
concl = shape(g5.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, 8420000, BAND1[0] - 40000, 1330000, BAND1[1] + 80000, RED, None, adj=0.25)
text_into(concl, [[('Cùng một khu vực', 7.5, 'FFFFFF', True)], [('xe phải vào 2 lần', 7.5, 'FFFFFF', True)]])
set_timing(s, [[g1], [g2], [g3], [g4], [g5]])

# ================= Slide 4: optimized =================
s = prs.slides[3]
clear(s)
pos = {n: OPT.index(n) + 1 for n in P}
g1 = bins(s, 'Bước 1 – Cùng 9 thùng', '8A94A6', {n: n for n in P})
label(g1.shapes, 7930000, BAND1[0] - 40000, 1820000, BAND1[1] + 80000, [[('Thứ tự mới', 7, INK, True)], [('1 → 3 → 4 → 7 → 8 → 9 → 6 → 5 → 2', 6.5, GREEN, True)]], GREEN)
g2 = group(s, 'Bước 2 – Cửa → #1 → #3 → #4 → #7')
poly(g2.shapes, route(chain([1, 3, 4, 7], True)), GREEN)
for n in (1, 3, 4, 7):
    x, y = P[n]
    b = shape(g2.shapes, MSO_SHAPE.OVAL, x + 35000, y - 175000, 170000, 170000, GREEN, 'FFFFFF', 1.0)
    text_into(b, [[(str(n), 7.5, 'FFFFFF', True)]], inset=0)
g3 = group(s, 'Bước 3 – #7 → #8 → #9 → #6')
poly(g3.shapes, route(chain([7, 8, 9, 6])), GREEN)
for n in (8, 9, 6):
    x, y = P[n]
    b = shape(g3.shapes, MSO_SHAPE.OVAL, x + 35000, y - 175000, 170000, 170000, GREEN, 'FFFFFF', 1.0)
    text_into(b, [[(str(n), 7.5, 'FFFFFF', True)]], inset=0)
ring(g3.shapes, 6, 9, GREEN)
label(g3.shapes, 8750000, BAND2[0], 1450000, BAND2[1], [[('#9 → #6 hạ liên tiếp', 7.5, GREEN, True)]], GREEN)
g4 = group(s, 'Bước 4 – #6 → #5 → #2 → về cửa')
pts = route(chain([6, 5, 2])) + door_back(P[2], 70000)[1:]
poly(g4.shapes, pts, GREEN)
for n in (5, 2):
    x, y = P[n]
    b = shape(g4.shapes, MSO_SHAPE.OVAL, x + 35000, y - 175000, 170000, 170000, GREEN, 'FFFFFF', 1.0)
    text_into(b, [[(str(n), 7.5, 'FFFFFF', True)]], inset=0)
ring(g4.shapes, 2, 5, GREEN)
label(g4.shapes, 7000000, BAND2[0], 1450000, BAND2[1], [[('#5 → #2 hạ liên tiếp', 7.5, GREEN, True)]], GREEN)
g5 = group(s, 'Bước 5 – Kết luận')
concl = shape(g5.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, 10340000, BAND1[0] - 40000, 1110000, BAND1[1] + 80000, GREEN, None, adj=0.25)
text_into(concl, [[('Mỗi khu vực', 7.5, 'FFFFFF', True)], [('chỉ vào 1 lần', 7.5, 'FFFFFF', True)]])
set_timing(s, [[g1], [g2], [g3], [g4], [g5]])

# ---- agreed wording for limitation 01 on every slide ----
NEW = {'Nearby bins missed across different aisles': 'Nearby bins across aisles not handled in sequence',
       'Không gom các vị trí gần nhau ở các dãy khác nhau': 'Điểm gần nhau khác dãy không được hạ hàng liên tiếp'}
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() in NEW:
                        r.text = NEW[r.text.strip()]

prs.save(DST)
print(f'saved {DST}  current {L_CUR / 1e5:.1f}  optimized {L_OPT / 1e5:.1f}  (-{SAVE}%)')
