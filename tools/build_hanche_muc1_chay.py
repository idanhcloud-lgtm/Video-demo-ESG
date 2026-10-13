"""Limitation 01 with a full TO on the map, everything but the focus area dimmed, and routes drawn as they run.

- The whole TO (13 bins) is shown; numbers follow the printed order (aisle sweep A → T, same rule as the app).
- Then the map outside the focus area (aisles C|D … D|…|…) is dimmed: #5 and #7 sit side by side.
- Routes are revealed segment by segment (wipe in the travel direction).
- Slide 3 adds a short green arrow #5 → #7: the two bins are right next to each other.
Slides: 3 = current (#5 → #6 → #7 → #8), 4 = algorithm (#5 → #7 → #6 → #8), 5 = both on one map.
Input: original deck with slide 4 duplicated after itself (scripts/add_slide.py deck slide4.xml --after slide4.xml).
Usage: python3 tools/build_hanche_muc1_chay.py <in.pptx> <out.pptx>
"""
import io
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

import build_hanche_muc1 as m

SRC, DST = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
m.load_carton(prs)
SW, SH = prs.slide_width, prs.slide_height

# ---- the TO: (aisle index, y). Aisles: 0 A|B, 1 B|C, 2 C|D, 3 D|…, 4 …|…, 6 Q|R, 7 R|S, 8 S|T ----
RAW = [(0, 4950000), (0, 3500000), (1, 3800000), (1, 5050000), (2, 5110000), (2, 2420000), (3, 5110000),
       (4, 2500000), (6, 2460000), (6, 3900000), (7, 5000000), (8, 3300000), (8, 4600000)]


def sweep(bins):
    """Current practice: aisles A → T, each swept one way from the end nearer the truck (front = bottom)."""
    out, cur = [], m.Y_FRONT
    for a in sorted({k for k, _ in bins}):
        g = sorted((y for k, y in bins if k == a), reverse=True)  # front (large y) first
        if abs(cur - g[0]) > abs(cur - g[-1]):
            g.reverse()
        out += [(a, y) for y in g]
        cur = out[-1][1]
    return out


TO = sweep(RAW)
P = {i + 1: (m.AISLE(a), y) for i, (a, y) in enumerate(TO)}
assert P[5][0] == m.AISLE(2) and P[5][1] > 5e6 and P[7][0] == m.AISLE(3), 'example no longer separates #5 and #7'
FOCUS = [5, 6, 7, 8]
CUR, OPT = [5, 6, 7, 8], [5, 7, 6, 8]
FX0, FX1 = m.AISLE(2) - 330000, m.AISLE(4) + 330000      # focus area (x)
FY0, FY1 = 2080000, 5560000                             # focus area (y)
MAP = (6197600, 1530000, 11480800, 6420000)             # map area inside the card
B1, B2 = m.BAND1, m.BAND2
SPEED = 550 / 1e6                                       # ms per EMU of route


def put_bins(sc, color, focus_color=None, dim=False):
    for n, (x, y) in P.items():
        f = n in FOCUS
        pic = sc.add_picture(io.BytesIO(m.CARTON), Emu(int(x - 200000)), Emu(int(y - 75000)), Emu(165100), Emu(149622))
        c = (focus_color or color) if f else color
        b = m.shape(sc, MSO_SHAPE.OVAL, x + 35000, y - 185000, 190000, 190000, c, 'FFFFFF', 1.0)
        m.text_into(b, [[(str(n), 7.5, 'FFFFFF', True)]], inset=0)


def dim_overlay(slide, name):
    """White veil over the map except the focus area (4 rectangles around it)."""
    g = m.group(slide, name)
    x0, y0, x1, y1 = MAP
    for (a, b, c, d) in [(x0, y0, x1, FY0), (x0, FY1, x1, y1), (x0, FY0, FX0, FY1), (FX1, FY0, x1, FY1)]:
        r = m.shape(g.shapes, MSO_SHAPE.RECTANGLE, a, b, c - a, d - b, 'FFFFFF', None)
        sf = r._element.spPr.find(qn('a:solidFill'))[0]
        etree.SubElement(sf, qn('a:alpha')).set('val', '62000')
    edge = m.shape(g.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, FX0, FY0, FX1 - FX0, FY1 - FY0, None, '8A94A6', 1.0, dash=MSO_LINE.DASH, adj=0.04)
    return g


def box(sc, x, w, band, lines, color, filled=False):
    s = m.shape(sc, MSO_SHAPE.ROUNDED_RECTANGLE, x, band[0] - 60000, w, band[1] + 120000, color if filled else 'FFFFFF', None if filled else color, 1.5, adj=0.25)
    m.text_into(s, lines)
    return s


def pair_ring(sc, color):
    (x1, y1), (x2, _) = P[5], P[7]
    return m.shape(sc, MSO_SHAPE.OVAL, x1 - 300000, y1 - 330000, x2 - x1 + 600000, 480000, None, color, 2.0, dash=MSO_LINE.DASH)


def running_route(slide, name, order, color, w=2.75, dash=None, off=0):
    """Route as straight segments, each revealed in its travel direction.
    Shapes stay top-level (PowerPoint cannot animate shapes inside a group). Returns the effects for the timing."""
    sc = slide.shapes
    legs = []
    for k, (a, b) in enumerate(zip(order, order[1:])):
        d = off + (k - 1) * 40000  # each leg slightly offset so up/down legs in one aisle stay visible
        legs.append([(x + d, y) for x, y in m.leg(P[a], P[b])])
    effects, t = [], 0
    for leg in legs:
        segs = [(p, q) for p, q in zip(leg, leg[1:]) if abs(q[0] - p[0]) + abs(q[1] - p[1]) >= 1]
        for i, (p, q) in enumerate(segs):
            L = abs(q[0] - p[0]) + abs(q[1] - p[1])
            s = m.poly(sc, [p, q], color, w, arrow=i == len(segs) - 1, dash=dash)  # arrowhead at each stop
            s.name = f'{name} – đoạn {len(effects) + 1}'
            d = 'up' if q[1] < p[1] else 'down' if q[1] > p[1] else 'right' if q[0] > p[0] else 'left'
            dur = max(150, int(L * SPEED))
            effects.append(('wipe', s, t, dur, d))
            t += dur
    return effects


WIPE = {'up': (4, 'wipe(up)'), 'down': (1, 'wipe(down)'), 'right': (8, 'wipe(right)'), 'left': (2, 'wipe(left)')}


def timing(slide, clicks):
    """clicks: list of lists of effects. ('fade', shape) | ('wipe', shape, delay, dur, dir) | ('appear', ...) | ('motion', shape, delay, dur, path)."""
    ids = iter(range(3, 100000))
    pars = []
    for click in clicks:
        inner = []
        for k, e in enumerate(click):
            kind, sh = e[0], e[1]
            delay = e[2] if len(e) > 2 else 0
            node = 'clickEffect' if k == 0 else 'withEffect'
            a, b, c = next(ids), next(ids), next(ids)
            sp = f'<p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl>'
            vis = (f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{sp}'
                   '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
            if kind == 'fade':
                body, head = vis + f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="500"/>{sp}</p:cBhvr></p:animEffect>', 'presetID="10" presetClass="entr" presetSubtype="0"'
            elif kind == 'wipe':
                sub, flt = WIPE[e[4]]
                body, head = vis + f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr><p:cTn id="{c}" dur="{e[3]}"/>{sp}</p:cBhvr></p:animEffect>', f'presetID="22" presetClass="entr" presetSubtype="{sub}"'
            elif kind == 'appear':
                body, head = vis, 'presetID="1" presetClass="entr" presetSubtype="0"'
            else:  # motion
                body = (f'<p:animMotion origin="layout" path="{e[4]}" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn id="{c}" dur="{e[3]}" fill="hold"/>{sp}'
                        '<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>')
                head = 'presetID="0" presetClass="path" presetSubtype="0" accel="0" decel="0"'
            inner.append(f'<p:par><p:cTn id="{a}" {head} fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')
        o1, o2 = next(ids), next(ids)
        pars.append(f'<p:par><p:cTn id="{o1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                    f'<p:par><p:cTn id="{o2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>' + ''.join(inner) +
                    '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{m.PNS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + ''.join(pars) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml)
    ext = slide._element.find(qn('p:extLst'))
    (ext.addprevious(el) if ext is not None else slide._element.append(el))


def prepare(slide):
    m.clear(slide)
    for sh in list(slide.shapes):
        if sh.name.startswith('Bước'):
            sh._element.getparent().remove(sh._element)


def fade(*groups):
    return [('fade', g) for g in groups]


NOTE = [[('Phiếu TO 13 điểm', 7.5, m.INK, True)], [('in theo thứ tự dãy A → T', 7, m.INK, False)]]
LX, LW = 9230000, 2220000   # labels sit right of the focus area

# ---------------- Slide 3: current ----------------
s = prs.slides[2]
prepare(s)
g1 = m.group(s, 'Bước 1 – Cả phiếu TO')
put_bins(g1.shapes, m.BLUE)
veil = dim_overlay(s, 'Bước 2 – Làm mờ, tập trung #5 và #7')
note = m.group(s, 'Bước 1 – Ghi chú phiếu')
box(note.shapes, 6280000, 1240000, B1, NOTE, m.BLUE)
g2 = m.group(s, 'Bước 2 – Cặp #5 · #7')
pair_ring(g2.shapes, m.RED)
box(g2.shapes, LX, LW, B2, [[('#5 và #7 nằm sát nhau (chỉ cách 1 lối)', 7.5, m.RED, True)], [('nhưng không hạ liên tiếp', 7, m.INK, False)]], m.RED)
run = running_route(s, 'Bước 3 – Xe chạy theo phiếu', CUR, m.RED)
lab = m.group(s, 'Bước 3 – Chú thích đường đi')
box(lab.shapes, LX, LW, B1, [[('Hạ #5 → lên hạ #6', 7.5, m.RED, True)], [('→ xuống hạ #7 → lại lên #8', 7.5, m.RED, True)]], m.RED)
(x5, y5), (x7, _) = P[5], P[7]
near = m.poly(s.shapes, [(x5 - 120000, y5 + 210000), (x7 + 60000, y5 + 210000)], m.GREEN, 5)
near.name = 'Bước 4 – Mũi tên xanh: #5 → #7 gần nhau'
g4 = m.group(s, 'Bước 4 – Kết luận')
c = m.shape(g4.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, LX, 5600000 - 240000, LW, 230000, m.RED, None, adj=0.3)
m.text_into(c, [[('Cùng một khu vực, xe phải vào 2 lần', 8, 'FFFFFF', True)]])
timing(s, [fade(g1, note), fade(veil, g2), run + [('fade', lab, 0)], [('wipe', near, 0, 500, 'right'), ('fade', g4, 300)]])

# ---------------- Slide 4: algorithm ----------------
s = prs.slides[3]
prepare(s)
g1 = m.group(s, 'Bước 1 – Cả phiếu TO')
put_bins(g1.shapes, '8A94A6', m.GREEN)
veil = dim_overlay(s, 'Bước 1 – Làm mờ')
note = m.group(s, 'Bước 1 – Ghi chú phiếu')
box(note.shapes, 6280000, 1240000, B1, [[('Cùng phiếu TO', 7.5, m.INK, True)], [('sắp lại thứ tự', 7, m.INK, False)]], '8A94A6')
run = running_route(s, 'Bước 2 – Xe chạy theo thuật toán', OPT, m.GREEN)
lab = m.group(s, 'Bước 2 – Chú thích đường đi')
box(lab.shapes, LX, LW, B1, [[('Hạ #5 → #7 liền nhau', 7.5, m.GREEN, True)], [('→ lên hạ #6 → đi tiếp #8', 7.5, m.GREEN, True)]], m.GREEN)
g3 = m.group(s, 'Bước 3 – Kết luận')
pair_ring(g3.shapes, m.GREEN)
box(g3.shapes, LX, LW, B2, [[('#5 và #7 hạ liên tiếp', 7.5, m.GREEN, True)]], m.GREEN)
c = m.shape(g3.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, LX, 5600000 - 240000, LW, 230000, m.GREEN, None, adj=0.3)
m.text_into(c, [[('Mỗi khu vực chỉ vào 1 lần', 8, 'FFFFFF', True)]])
timing(s, [fade(g1, veil, note), run + [('fade', lab, 0)], fade(g3)])

# ---------------- Slide 5: both ----------------
s = prs.slides[4]
prepare(s)
for sh in s.shapes:
    if sh.has_text_frame and 'Optimized Route' in sh.text_frame.text:
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        runs[0].text = 'Route Comparison'
        for r in runs[1:]:
            r.text = ''
        if len(runs) > 1:
            runs[1].text = ' / So sánh: theo phiếu A → T và thuật toán'
g1 = m.group(s, 'Bước 1 – Cả phiếu TO + chú giải')
put_bins(g1.shapes, m.BLUE)
box(g1.shapes, LX, LW, B1, [[('- - -  ', 8, m.RED, True), ('Theo phiếu: #5 → #6 → #7 → #8', 7, m.INK, True)], [('——  ', 8, m.GREEN, True), ('Thuật toán: #5 → #7 → #6 → #8', 7, m.INK, True)]], '8A94A6')
veil = dim_overlay(s, 'Bước 1 – Làm mờ')
run_r = running_route(s, 'Bước 2 – Theo phiếu', CUR, m.RED, 2.0, MSO_LINE.DASH, off=-45000)
run_g = running_route(s, 'Bước 3 – Thuật toán', OPT, m.GREEN, 2.75, off=45000)
g4 = m.group(s, 'Bước 4 – Kết luận')
pair_ring(g4.shapes, m.NAVY)
box(g4.shapes, LX, LW, B2, [[('Theo phiếu: #5 và #7 cách nhau', 7, m.RED, True)], [('Thuật toán: #5 và #7 liền nhau', 7, m.GREEN, True)]], m.NAVY)
timing(s, [fade(g1, veil), run_r, run_g, fade(g4)])

prs.save(DST)
print('saved', DST, ' TO order:', TO)
