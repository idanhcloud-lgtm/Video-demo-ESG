"""Limitation 01 on one slide: a virtual TO on the deck's map, both full routes, then only the limitation lit up.

- Bins: a virtual TO of 13 bins on the schematic map (aisles between the drawn rack columns).
- Red  = current practice (sweepOrder in the app): aisles A → T, each aisle swept one way from the nearer end.
- Green = optimized (as the app): serpentine start, then 2-opt on the total distance.
  Distances: rectilinear, changing aisle at the nearest crossing (front, 2 cross aisles, far end); start/end at the dock.
- Last click dims the map except the nearby pair the current order splits up (found automatically).
Usage: python3 tools/build_hanche_toan_map.py <in.pptx> <out.pptx>   (slide 3 of the original 'Hạn chế' deck is rebuilt)
"""
import io
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

import build_hanche_muc1 as m

SRC, DST = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
m.load_carton(prs)

# ---- virtual TO: (aisle index, y). Aisles: 0 A|B, 1 B|C, 2 C|D, 3 D|…, 4 …|…, 5 …|Q, 6 Q|R, 7 R|S, 8 S|T ----
RAW = [(0, 4950000), (0, 3500000), (1, 3800000), (1, 5050000), (2, 5110000), (2, 2420000), (3, 5110000),
       (4, 2500000), (6, 2460000), (6, 3900000), (7, 5000000), (8, 3300000), (8, 4600000)]
CROSS = [m.Y_TOP, m.Y_C1, m.Y_C2, m.Y_FRONT]
DOOR = (m.DOOR_X, m.Y_FRONT)


def dist(p, q):
    (x1, y1), (x2, y2) = p, q
    if abs(x1 - x2) < 1:
        return abs(y1 - y2)
    return abs(x1 - x2) + min(abs(y1 - c) + abs(y2 - c) for c in CROSS)


def door_dist(p):
    return abs(p[0] - DOOR[0]) + abs(p[1] - DOOR[1])


def length(seq):
    return door_dist(P[seq[0]]) + sum(dist(P[a], P[b]) for a, b in zip(seq, seq[1:])) + door_dist(P[seq[-1]])


def sweep(bins):
    out, cur = [], m.Y_FRONT
    for a in sorted({k for k, _ in bins}):
        g = sorted((y for k, y in bins if k == a), reverse=True)  # front (large y) first
        if abs(cur - g[0]) > abs(cur - g[-1]):
            g.reverse()
        out += [(a, y) for y in g]
        cur = out[-1][1]
    return out


TO = sweep(RAW)                                   # bin numbers = printed order
P = {i + 1: (m.AISLE(a), y) for i, (a, y) in enumerate(TO)}
CUR = list(P)
# optimized: serpentine start (even aisles from the front, odd aisles from the far end), then 2-opt
seq = sorted(P, key=lambda n: (TO[n - 1][0], -TO[n - 1][1] if TO[n - 1][0] % 2 == 0 else TO[n - 1][1]))
best, improved = length(seq), True
while improved:
    improved = False
    for i in range(len(seq) - 1):
        for j in range(i + 1, len(seq)):
            cand = seq[:i] + seq[i:j + 1][::-1] + seq[j + 1:]
            if length(cand) < best - 1:
                seq, best, improved = cand, length(cand), True
OPT = seq
SAVE = round(100 * (length(CUR) - length(OPT)) / length(CUR))
# the limitation: nearest pair that the algorithm handles back to back but the printed order splits up
pos = {n: i for i, n in enumerate(OPT)}
PAIR = min(((a, b) for a in P for b in P if a < b and b - a > 1 and abs(pos[a] - pos[b]) == 1), key=lambda ab: dist(P[ab[0]], P[ab[1]]))
A, Bn = PAIR

SPEED_MS = 200 / 1e6   # ms per EMU of route
B1, B2 = m.BAND1, m.BAND2
LX, LW = 9230000, 2220000


def full_path(order, off):
    legs = [m.door_out(P[order[0]])] + [m.leg(P[a], P[b]) for a, b in zip(order, order[1:])] + [m.door_back(P[order[-1]], 60000)]
    return [[(x + off, y + off) for x, y in leg] for leg in legs]


def draw_route(slide, name, order, color, w, dash, off):
    effects, t = [], 0
    for leg in full_path(order, off):
        segs = [(p, q) for p, q in zip(leg, leg[1:]) if abs(q[0] - p[0]) + abs(q[1] - p[1]) >= 1]
        for i, (p, q) in enumerate(segs):
            L = abs(q[0] - p[0]) + abs(q[1] - p[1])
            s = m.poly(slide.shapes, [p, q], color, w, arrow=i == len(segs) - 1, dash=dash)
            s.name = f'{name} – đoạn {len(effects) + 1}'
            d = 'up' if q[1] < p[1] else 'down' if q[1] > p[1] else 'right' if q[0] > p[0] else 'left'
            dur = max(80, int(L * SPEED_MS))
            effects.append(('wipe', s, t, dur, d))
            t += dur
    return effects


def box(sc, x, w, band, lines, color, filled=False, grow=0):
    s = m.shape(sc, MSO_SHAPE.ROUNDED_RECTANGLE, x, band[0] - 60000 - grow, w, band[1] + 120000 + 2 * grow, color if filled else 'FFFFFF', None if filled else color, 1.5, adj=0.2)
    m.text_into(s, lines)
    return s


def veil(slide, name, fx0, fy0, fx1, fy1):
    g = m.group(slide, name)
    x0, y0, x1, y1 = 6197600, 1530000, 11480800, 6420000
    for a, b, c, d in [(x0, y0, x1, fy0), (x0, fy1, x1, y1), (x0, fy0, fx0, fy1), (fx1, fy0, x1, fy1)]:
        r = m.shape(g.shapes, MSO_SHAPE.RECTANGLE, a, b, c - a, d - b, 'FFFFFF', None)
        etree.SubElement(r._element.spPr.find(qn('a:solidFill'))[0], qn('a:alpha')).set('val', '70000')
    return g


s = prs.slides[2]
m.clear(s)

# step 1: bins (numbers = printed order) + legend
g1 = m.group(s, 'Bước 1 – Phiếu TO ảo 13 điểm')
for n, (x, y) in P.items():
    g1.shapes.add_picture(io.BytesIO(m.CARTON), Emu(int(x - 200000)), Emu(int(y - 75000)), Emu(165100), Emu(149622))
    b = m.shape(g1.shapes, MSO_SHAPE.OVAL, x + 35000, y - 185000, 190000, 190000, m.BLUE, 'FFFFFF', 1.0)
    m.text_into(b, [[(str(n), 7.5, 'FFFFFF', True)]], inset=0)


def chip(name, x, w, lines, color):
    """Legend chip in the '35 dãy kệ' row above the racks (over the decorative arrows)."""
    g = m.group(s, name)
    c = m.shape(g.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1765000, w, 200000, 'FFFFFF', color, 1.25, adj=0.5)
    m.text_into(c, lines, inset=20000)
    return g


# steps 2, 3: full routes
red = draw_route(s, 'Bước 2 – Theo dãy A → T', CUR, m.RED, 2.0, MSO_LINE.DASH, -45000)
lab_r = chip('Bước 2 – Chú giải đỏ', 6230000, 1900000, [[('- - -  Theo dãy A → T', 7, m.RED, True), ('  (số = thứ tự in)', 6.5, m.INK, False)]], m.RED)
green = draw_route(s, 'Bước 3 – Tối ưu', OPT, m.GREEN, 2.75, None, 45000)
lab_g = chip('Bước 3 – Chú giải xanh', 9550000, 1880000, [[('——  Tối ưu (thuật toán)', 7, m.GREEN, True)]], m.GREEN)

# step 4: dim everything but the pair the printed order splits up
(xa, ya), (xb, yb) = P[A], P[Bn]
fx0, fx1 = min(xa, xb) - 330000, max(xa, xb) + 330000
fy0, fy1 = min(ya, yb) - 380000, max(ya, yb) + 330000
v = veil(s, 'Bước 4 – Làm mờ, chỉ sáng điểm hạn chế', fx0, fy0, fx1, fy1)
g4 = m.group(s, 'Bước 4 – Điểm hạn chế')
m.shape(g4.shapes, MSO_SHAPE.OVAL, fx0, fy0, fx1 - fx0, fy1 - fy0, None, m.RED, 2.25, dash=MSO_LINE.DASH)
box(g4.shapes, LX, LW, (4750000, 250000),
    [[(f'#{A} và #{Bn} nằm sát nhau', 7.5, m.RED, True)],
     [(f'Theo dãy: cách nhau {Bn - A - 1} điểm, xe vào 2 lần', 6.5, m.RED, True)],
     [('Tối ưu: hạ liên tiếp', 6.5, m.GREEN, True)]], m.RED, grow=60000)
res = m.shape(g4.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, LX, 5330000, LW, 230000, m.GREEN, None, adj=0.3)
m.text_into(res, [[(f'Lộ trình tối ưu ngắn hơn khoảng {SAVE}%', 7.5, 'FFFFFF', True)]])
# labels stay readable above the veil
for g in (lab_r, lab_g):
    s.shapes._spTree.remove(g._element)
    s.shapes._spTree.append(g._element)

m.timing_fx(s, [[('fade', g1)], red + [('fade', lab_r, 0)], green + [('fade', lab_g, 0)], [('fade', v), ('fade', g4)]])
prs.save(DST)
print(f'saved {DST}\n current {CUR} {length(CUR) / 1e5:.0f}\n optimized {OPT} {length(OPT) / 1e5:.0f} (-{SAVE}%)\n pair {PAIR}')
