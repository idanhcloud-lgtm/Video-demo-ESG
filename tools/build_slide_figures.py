"""Rebuild the right-hand figure of slides 4-9 (limitations) with native PowerPoint shapes.
# Usage: python3 tools/build_slide_figures.py <in.pptx> <out.pptx> <geom.json>
# geom.json: map projection + current route polyline (pts, stopIdx) exported from src/map/route.ts.

Only the current process (problems) is shown - no algorithm/solution. Orientation follows the
demo video: dispatch doors on the left, receiving at the far end on the right.
Every step is a group that fades in on click.
"""
import copy, io, json, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

SRC, DST, GEOM = sys.argv[1], sys.argv[2], sys.argv[3]
G = json.load(open(GEOM))
prs = Presentation(SRC)

# ---------- map geometry (same projection as the video) ----------
VT0, PXVT, LANE1, LANEPX = 103.2, 7.317, 596.3, 32.12
vtX = lambda v: VT0 + v * PXVT
laneY = lambda l: LANE1 - (l - 1) * LANEPX
RACKS = [("CAT",1,"T",76),("CBP",1,"P",76),("CBT",2,"T",76),("CCP",2,"P",76),("CCT",3,"T",80),("CDP",3,"P",80),
 ("CDT",4,"T",104),("CEP",4,"P",104),("CET",5,"T",104),("CFP",5,"P",104),("CFT",6,"T",104),("CGP",6,"P",104),
 ("CGT",7,"T",104),("CHP",7,"P",104),("CHT",8,"T",104),("CIP",8,"P",104),("CIT",9,"T",104),("CKP",9,"P",104),
 ("CKT",10,"T",104),("CLP",10,"P",104),("CLT",11,"T",104),("CMP",11,"P",104),("CMT",12,"T",104),("CNP",12,"P",104),
 ("CNT",13,"T",104),("COP",13,"P",104),("COT",14,"T",104),("CPP",14,"P",104),("CPT",15,"T",104),("CQP",15,"P",104),
 ("CQT",16,"T",104),("CRP",16,"P",104),("CRT",17,"T",104),("CSP",17,"P",104),("CTP",18,"P",96)]
LANE_OF = {c: l for c, l, _, _ in RACKS}
TO28 = [("CAT",18),("CAT",28),("CBT",55),("CBT",68),("CCT",70),("CDP",20),("CDT",12),("CDT",62),("CET",38),("CET",98),
 ("CGT",85),("CHP",76),("CHT",95),("CIP",24),("CIT",34),("CIT",50),("CLT",45),("CLT",60),("CMT",88),("CMT",100),
 ("COT",28),("CPP",15),("CPT",72),("CPT",82),("CQT",38),("CRT",52),("CTP",60),("CTP",92)]
STOP = [(vtX(v), laneY(LANE_OF[r])) for r, v in TO28]          # current visit order = TO28 order (A -> T)
CUR, SIDX = G["cur"]["pts"], G["cur"]["stopIdx"]
def legs(a, b):  # polyline from stop a to stop b (1-based, current order)
    return CUR[SIDX[a - 1]: SIDX[b - 1] + 1]

# ---------- placement inside the map card ----------
CARD = (5918200, 1625600, 5816600, 4826000)
PAD = 110000
HEAD = 300000                       # title row above the map
VW, VH = 962.0, 628.0               # view box in map px
S = (CARD[2] - 2 * PAD) / VW        # EMU per map px
OX, OY = CARD[0] + PAD, CARD[1] + PAD + HEAD
MAP_BOTTOM = OY + VH * S
X = lambda px: int(OX + px * S)
Y = lambda py: int(OY + py * S)
L = lambda px: int(px * S)

NAVY, DISP, GREENB = "0B2470", "0B3E9A", "03A176"
RED, RED_SOFT, AMBER, INK, MUTED = "D7263D", "FBE3E6", "E8A317", "1F2A44", "6B7280"
RACK, FLOOR, RECV = "D9C59C", "EEF2F7", "E4EEF9"

def rgb(h): return RGBColor.from_string(h)

def fill(sh, color=None):
    if color is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color)

def line(sh, color=None, w=0.75, dash=None):
    if color is None:
        sh.line.fill.background(); return
    sh.line.color.rgb = rgb(color); sh.line.width = Pt(w)
    if dash: sh.line.dash_style = dash

def no_shadow(sh):
    sp = sh._element.spPr
    if sp.find(qn("a:effectLst")) is None:
        sp.append(etree.SubElement(sp, qn("a:effectLst")))

def rect(sh_coll, x, y, w, h, color, ln=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, name=None, dash=None, adj=None):
    s = sh_coll.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(max(int(w), 1)), Emu(max(int(h), 1)))
    fill(s, color); line(s, ln, lw, dash); no_shadow(s)
    if adj is not None: s.adjustments[0] = adj
    if name: s.name = name
    s.text_frame.text = ""
    return s

def text(sh_coll, x, y, w, h, runs, size=8, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         rot=0, name=None, box=None, ln=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, adj=None, inset=0):
    """runs: str or list of lines; each line str or list of (text, size, color, bold)."""
    if box is None and ln is None:
        s = sh_coll.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    else:
        s = rect(sh_coll, x, y, w, h, box, ln, lw, shape=shape, adj=adj)
    tf = s.text_frame
    tf.word_wrap = True; tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(inset))
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [runs]
    for i, ln_ in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = ln_ if isinstance(ln_, list) else [(ln_, size, color, bold)]
        for t, sz, c, b in parts:
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(sz); f.bold = b; f.color.rgb = rgb(c); f.name = "Arial"
    if rot: s.rotation = rot
    if name: s.name = name
    return s

def poly(sh_coll, pts, color, w=2.0, dash=None, arrow=True, name=None, trim=13):
    pts = [tuple(p) for p in pts]
    if arrow and trim and len(pts) >= 2:  # stop the arrowhead at the edge of the target dot
        (ax, ay), (bx, by) = pts[-2], pts[-1]
        d = ((bx - ax) ** 2 + (by - ay) ** 2) ** 0.5
        if d > trim * 1.5: pts[-1] = (bx - (bx - ax) * trim / d, by - (by - ay) * trim / d)
    fb = sh_coll.build_freeform(X(pts[0][0]), Y(pts[0][1]), scale=1.0)
    fb.add_line_segments([(X(px), Y(py)) for px, py in pts[1:]], close=False)
    s = fb.convert_to_shape()
    fill(s, None); line(s, color, w, dash); no_shadow(s)
    ln = s._element.spPr.find(qn("a:ln"))
    etree.SubElement(ln, qn("a:round"))
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    if name: s.name = name
    return s

def dot(sh_coll, px, py, r, color, n=None, ln="FFFFFF", fsize=6.5, fcolor="FFFFFF"):
    s = rect(sh_coll, X(px) - L(r), Y(py) - L(r), L(2 * r), L(2 * r), color, ln, 0.75, shape=MSO_SHAPE.OVAL)
    if n is not None:
        tf = s.text_frame; tf.word_wrap = False
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, side, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r_ = p.add_run(); r_.text = str(n)
        r_.font.size = Pt(fsize); r_.font.bold = True; r_.font.color.rgb = rgb(fcolor); r_.font.name = "Arial"
    return s

def group(slide, name):
    g = slide.shapes.add_group_shape(); g.name = name
    return g

# ---------- base map ----------
DOOR_Y = [d["y"] for d in G["doors"]]

def base_map(slide, title, legend=True):
    g = group(slide, "Bản đồ Kho C")
    sc = g.shapes
    # title row
    text(sc, CARD[0] + PAD, CARD[1] + PAD - 20000, 3600000, 260000,
         [[("Kho C – Duy Tân", 10, NAVY, True), ("  ·  " + title, 9, MUTED, False)]], align=PP_ALIGN.LEFT)
    if legend:
        lx = CARD[0] + CARD[2] - PAD - 2150000
        s = sc.add_connector(1, Emu(lx), Emu(CARD[1] + PAD + 110000), Emu(lx + 260000), Emu(CARD[1] + PAD + 110000))
        s.line.color.rgb = rgb(RED); s.line.width = Pt(2.25)
        text(sc, lx + 300000, CARD[1] + PAD - 20000, 1850000, 260000, "Hiện tại: hạ hàng theo dãy A → T", 8, RED, True, align=PP_ALIGN.LEFT)
    # floor, dispatch, receiving
    rect(sc, X(80), Y(20), L(808), L(600), FLOOR, shape=MSO_SHAPE.RECTANGLE)
    rect(sc, X(36), Y(4), L(40), L(620), DISP, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
    text(sc, X(56) - L(300), Y(314) - L(13), L(600), L(26), "KHU XUẤT HÀNG · DISPATCH", 7.5, "FFFFFF", True, rot=-90)
    for i, y in enumerate(DOOR_Y):
        rect(sc, X(4), Y(y - 7.5), L(27), L(15), "FFFFFF", DISP, 0.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
    for i in (0, 12, 24):
        text(sc, X(4), Y(DOOR_Y[i] - 7.5), L(27), L(15), f"D{i + 1}", 5, DISP, True)
    rect(sc, X(890), Y(4), L(66), L(620), RECV, "5B8DD6", 0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12, dash=MSO_LINE.DASH)
    text(sc, X(923) - L(300), Y(314) - L(13), L(600), L(26), "KHU NHẬP HÀNG · RECEIVING", 7.5, "2F5597", True, rot=90)
    # cross aisles (tunnels VT 33-35, 74-76)
    for vt in (34, 75):
        rect(sc, X(vtX(vt - 1.5)), Y(26), L(3 * PXVT), L(590), "FFF3D6")
        text(sc, X(vtX(vt)) - L(70), Y(6), L(140), L(18), "Lối ngang", 6.5, "A86A00", True)
    # racks
    for code, lane, side, vmax in RACKS:
        y = laneY(lane) + 8.3 if side == "T" else laneY(lane) - 8.3 - 6.6
        segs, start = [], -1
        for v in range(1, vmax + 2):
            ok = v <= vmax and not (33 <= v <= 35 or 74 <= v <= 76) and not (code == "CAT" and 9 <= v <= 14)
            if ok and start < 0: start = v
            if not ok and start >= 0: segs.append((start, v - 1)); start = -1
        for a, b in segs:
            rect(sc, X(vtX(a) - 3.6), Y(y - 1.2), L(vtX(b) - vtX(a) + 7.2), L(9.0), RACK)
    # rack letters (one per physical rack A..T) at the far end
    letters = {}
    for code, lane, side, _ in RACKS:
        y = laneY(lane) + 8.3 + 3.3 if side == "T" else laneY(lane) - 8.3 - 3.3
        letters.setdefault(code[1], []).append(y)
    for k, ys in letters.items():
        yc = sum(ys) / len(ys)
        text(sc, X(869), Y(yc - 7), L(19), L(14), k, 6, "8A6D2F", True)
    return g

# ---------- callouts ----------
def callout(sc, cx, cy, w, h, title, sub=None, color=RED):
    lines = [[(title, 8.5, color, True)]]
    if sub: lines.append([(sub, 7, INK, False)])
    return text(sc, X(cx) - L(w) / 2, Y(cy) - L(h) / 2, L(w), L(h), lines, box="FFFFFF", ln=color, lw=1.25,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.18, inset=40000)

def bottom_banner(slide, txt, sub=None, name="Kết luận"):
    g = group(slide, name)
    y0 = int(MAP_BOTTOM + 70000)
    h = CARD[1] + CARD[3] - PAD + 20000 - y0
    lines = [[(txt, 10.5, "FFFFFF", True)]]
    if sub: lines.append([(sub, 8.5, "FFE4E8", False)])
    text(g.shapes, CARD[0] + PAD, y0, CARD[2] - 2 * PAD, h, lines, box=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2, inset=50000)
    return g

def stops_dim(slide, skip=()):
    g = group(slide, "28 điểm hạ hàng (mờ)")
    for i, (x, y) in enumerate(STOP, 1):
        if i not in skip: dot(g.shapes, x, y, 5.5, "B9C2D0", ln="FFFFFF")
    return g

def stops_full(slide, name="28 điểm hạ hàng"):
    g = group(slide, name)
    for x, y in STOP: dot(g.shapes, x, y, 8, AMBER, ln="7A5200")
    return g

def numbered(sc, ns, color=RED):
    for n in ns:
        x, y = STOP[n - 1]; dot(sc, x, y, 12.5, color, n=n)

# ---------- forklift picture (taken from the original slide 5) ----------
FORK = None
for sh in prs.slides[4].shapes:
    if sh.name == "!!Route 1 forklift 1": FORK = sh.image.blob
assert FORK

def forklift(sc, px, py, w=30, flip=False):
    h = w * 174090 / 254000
    p = sc.add_picture(io.BytesIO(FORK), Emu(X(px - w / 2)), Emu(Y(py - h / 2)), Emu(L(w)), Emu(L(h)))
    if flip:
        p._element.spPr.find(qn("a:xfrm")).set("flipH", "1")
    return p

# ---------- slide-level helpers ----------
def clear_map(slide):
    cx0, cy0 = CARD[0], CARD[1]; cx1, cy1 = cx0 + CARD[2], cy0 + CARD[3]
    for sh in list(slide.shapes):
        if sh.name.startswith("!!Map card"): continue
        if sh.left is None: continue
        mx, my = sh.left + sh.width / 2, sh.top + sh.height / 2
        if cx0 <= mx <= cx1 and cy0 <= my <= cy1:
            sh._element.getparent().remove(sh._element)
    t = slide._element.find(qn("p:timing"))
    if t is not None: slide._element.remove(t)

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
def set_timing(slide, clicks):
    """clicks: list of steps; each step is a list of (shape, delay_ms). First item starts on click."""
    ids = iter(range(3, 100000))
    pars = []
    for step in clicks:
        inner = []
        for k, (sh, delay) in enumerate(step):
            node = "clickEffect" if k == 0 else "withEffect"
            a, b, c = next(ids), next(ids), next(ids)
            inner.append(f'''<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="500"/><p:tgtEl><p:spTgt spid="{sh.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>''')
        o1, o2 = next(ids), next(ids)
        pars.append(f'<p:par><p:cTn id="{o1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                    f'<p:par><p:cTn id="{o2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                    + "".join(inner) + '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           + "".join(pars) +
           '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml)
    root = slide._element
    ext = root.find(qn("p:extLst"))
    if ext is not None: ext.addprevious(el)
    else: root.append(el)

# =====================================================================
slides = prs.slides
for i in range(3, 9): clear_map(slides[i])

# ---- Slide 4: overview ----
s = slides[3]
base_map(s, "Ví dụ 1 phiếu TO B2C, 28 điểm hạ hàng", legend=False)
g_dots = stops_full(s)
g_note = group(s, "Ghi chú 28 điểm"); callout(g_note.shapes, 470, 616 + 0, 330, 20, "1 phiếu TO B2C · 28 điểm hạ hàng", color=NAVY)
# move the note under the map
for sh in g_note.shapes: sh.top = int(MAP_BOTTOM + 90000); sh.height = 300000
BADGE = {1: (vtX(94), laneY(13.5)), 2: (vtX(45), laneY(3.5)), 3: (vtX(52), laneY(7.5)), 4: (vtX(5), laneY(1.55)), 5: (vtX(99), laneY(10))}
badges = []
for n, (bx, by) in BADGE.items():
    g = group(s, f"Hạn chế {n:02d}")
    rect(g.shapes, X(bx) - L(19), Y(by) - L(19), L(38), L(38), RED_SOFT, RED, 1.0, shape=MSO_SHAPE.OVAL)
    dot(g.shapes, bx, by, 13.5, NAVY, n=f"{n:02d}", fsize=7.5)
    badges.append(g)
set_timing(s, [[(g_dots, 0), (g_note, 0)]] + [[(g, 0)] for g in badges])

# ---- Slide 5: 01 nearby bins not in sequence ----
s = slides[4]
base_map(s, "01 · Điểm gần nhau khác dãy")
stops_dim(s, skip=(19, 20, 21, 22, 23, 24))
g1 = group(s, "Bước 1 – Cụm điểm gần nhau")
x0, x1 = vtX(68), vtX(103.5); y0, y1 = laneY(15.55), laneY(11.45)
rect(g1.shapes, X(x0), Y(y0), L(x1 - x0), L(y1 - y0), None, RED, 1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1, dash=MSO_LINE.DASH)
numbered(g1.shapes, [19, 20, 23, 24], AMBER)
callout(g1.shapes, vtX(86), laneY(9.9), 300, 36, "4 điểm sát nhau", "dãy M (#19, #20) và dãy P (#23, #24)")
g2 = group(s, "Bước 2 – Hạ #19, #20")
poly(g2.shapes, legs(19, 20), RED, 2.25); numbered(g2.shapes, [19, 20])
g3 = group(s, "Bước 3 – Đi xa hạ #21, #22")
poly(g3.shapes, legs(20, 22), RED, 2.25); numbered(g3.shapes, [19, 20, 21, 22])
callout(g3.shapes, vtX(20), laneY(16.9), 250, 36, "Đi xa về đầu kho", "hạ #21, #22 theo thứ tự dãy")
g4 = group(s, "Bước 4 – Quay lại hạ #23, #24")
poly(g4.shapes, legs(22, 24), RED, 2.25); numbered(g4.shapes, [21, 22, 23, 24])
callout(g4.shapes, vtX(58), laneY(16.6), 250, 36, "Quay lại khu vực cũ", "mới hạ #23, #24")
g5 = bottom_banner(s, "Cùng một khu vực, xe phải vào 2 lần", "Phiếu in theo dãy A → T nên các điểm gần nhau không được hạ liên tiếp")
set_timing(s, [[(g1, 0)], [(g2, 0)], [(g3, 0)], [(g4, 0)], [(g5, 0)]])

# ---- Slide 6: 02 cross aisles not optimally used ----
s = slides[5]
base_map(s, "02 · Lối ngang")
stops_dim(s, skip=(5, 6, 7, 8))
g1 = group(s, "Bước 1 – Điểm #5 đến #8")
numbered(g1.shapes, [5, 6, 7, 8], AMBER)
callout(g1.shapes, vtX(50), laneY(7.2), 240, 36, "Dãy C, D", "#5 → #6 → #7 → #8 theo phiếu")
g2 = group(s, "Bước 2 – Chạy ngược về đầu dãy")
poly(g2.shapes, legs(5, 7), RED, 2.25); numbered(g2.shapes, [5, 6, 7])
callout(g2.shapes, vtX(18), laneY(5.6), 250, 36, "Chạy ngược về đầu dãy", "để đổi lối")
g3 = group(s, "Bước 3 – Lại vào sâu")
poly(g3.shapes, legs(7, 8), RED, 2.25); numbered(g3.shapes, [8])
g4 = group(s, "Bước 4 – Lối ngang không được dùng")
for vt in (34, 75):
    rect(g4.shapes, X(vtX(vt - 1.5)) - L(2), Y(26), L(3 * PXVT + 4), L(590), None, AMBER, 1.5, dash=MSO_LINE.DASH)
callout(g4.shapes, vtX(80), laneY(5.4), 250, 36, "Lối ngang ngay cạnh", "nhưng xe không đi qua", color="A86A00")
g5 = bottom_banner(s, "Thứ tự trên phiếu quyết định chỗ đổi lối", "Xe chạy về đầu dãy rồi lại vào sâu, thay vì đi qua lối ngang gần nhất")
set_timing(s, [[(g1, 0)], [(g2, 0)], [(g3, 0)], [(g4, 0)], [(g5, 0)]])

# ---- Slide 7: 03 items scattered across many aisles ----
s = slides[6]
base_map(s, "03 · Hàng rải rác nhiều dãy")
g1 = stops_full(s)
g2 = group(s, "Bước 2 – Phân bố theo vùng")
for a, b, lab, n in ((0, 33, "Đầu kho", 7), (34, 74, "Giữa kho", 13), (75, 104, "Cuối kho", 8)):
    cx = (vtX(a) + vtX(b)) / 2
    cx = {0: vtX(15), 34: vtX(54.5), 75: vtX(93)}[a]
    text(g2.shapes, X(cx) - L(70), Y(5), L(140), L(17), [[(f"{lab}: ", 7, INK, True), (f"{n} điểm", 7, RED, True)]], box="FFFFFF", ln=RED, lw=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.4)
note = group(s, "Ghi chú 15/18 lối")
text(note.shapes, CARD[0] + PAD, int(MAP_BOTTOM + 70000), CARD[2] - 2 * PAD, CARD[1] + CARD[3] - PAD + 20000 - int(MAP_BOTTOM + 70000), [[("28 điểm nằm trên 15/18 lối, trải từ đầu kho tới cuối kho", 10.5, "FFFFFF", True)]], box=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2, inset=50000)
g3 = group(s, "Bước 3 – Lộ trình hiện tại")
poly(g3.shapes, CUR, RED, 1.5, arrow=False)
g4 = bottom_banner(s, "Không có lộ trình tối ưu, xe chạy khắp kho", "Hàng càng rải rác, đi theo thứ tự dãy A → T càng chạy thừa nhiều")
set_timing(s, [[(g1, 0)], [(g2, 0), (note, 0)], [(g3, 0)], [(g4, 0)]])
# keep dots above the route
g1._element.getparent().append(g1._element)
for gg in (note, g4): gg._element.getparent().append(gg._element)

# ---- Slide 8: 04 morning congestion in aisles A, B ----
s = slides[7]
base_map(s, "04 · Giờ cao điểm 9:00")
g1 = group(s, "Bước 1 – 9:00 sáng")
text(g1.shapes, X(vtX(72)), Y(laneY(17.2)), L(200), L(64), [[("9:00 AM", 18, "FFFFFF", True)], [("Giờ cao điểm sáng", 8, "FFFFFF", False)]],
     box=RED, ln="FFFFFF", lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
g2 = group(s, "Bước 2 – Các TO đều bắt đầu từ dãy A, B")
for l in (1, 2):
    rect(g2.shapes, X(vtX(-1)), Y(laneY(l) - 8), L(vtX(77) - vtX(-1)), L(16), "F7B4BD", shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
for k, l in (("A", 1.3), ("B", 2.0)):
    pass
callout(g2.shapes, vtX(80), laneY(11), 330, 36, "Các TO đều bắt đầu từ dãy A, B", "xe nào cũng vào dãy A, B trước")
g3 = group(s, "Bước 3 – Xe dồn vào dãy A, B")
for l in (1, 2):
    poly(g3.shapes, [(90, laneY(l)), (vtX(30), laneY(l))], RED, 2.0)
turn = [(4, 1), (6, 2), (8, 1)]
for l, src in turn:
    poly(g3.shapes, [(vtX(18), laneY(src)), (vtX(34), laneY(src)), (vtX(34), laneY(l)), (vtX(46), laneY(l))], "5B6B85", 1.5, dash=MSO_LINE.DASH)
    forklift(g3.shapes, vtX(50), laneY(l), 34)
for l in (1, 2):
    for v in (40, 47, 54):
        forklift(g3.shapes, vtX(v), laneY(l), 34)
g4 = group(s, "Bước 4 – 6 xe kẹt")
for l in (1, 2):
    for v in (40, 47, 54):
        rect(g4.shapes, X(vtX(v)) - L(17), Y(laneY(l)) - L(17), L(34), L(34), None, RED, 1.5, shape=MSO_SHAPE.OVAL)
callout(g4.shapes, vtX(81), laneY(3.6), 230, 36, "6 xe kẹt giữa dãy A, B", "chỉ 3 xe kịp rẽ lối ngang")
g5 = bottom_banner(s, "Ùn tắc ở dãy A, B vào đầu giờ sáng", "Mọi phiếu TO đều in bắt đầu từ dãy A, B nên xe dồn về cùng một chỗ")
set_timing(s, [[(g1, 0)], [(g2, 0)], [(g3, 0)], [(g4, 0)], [(g5, 0)]])

# ---- Slide 9: 05 storage by product group ----
s = slides[8]
base_map(s, "05 · Vị trí lưu kho", legend=False)
g1 = group(s, "Bước 1 – Vùng gần cửa / vùng sâu")
rect(g1.shapes, X(vtX(0.5)), Y(26), L(vtX(32.5) - vtX(0.5)), L(590), None, GREENB, 1.25, dash=MSO_LINE.DASH)
text(g1.shapes, X(vtX(0.5)), Y(598), L(vtX(32.5) - vtX(0.5)), L(18), "Gần cửa xuất", 7, GREENB, True, box="FFFFFF")
rect(g1.shapes, X(vtX(76.5)), Y(26), L(vtX(104.5) - vtX(76.5)), L(590), None, RED, 1.25, dash=MSO_LINE.DASH)
text(g1.shapes, X(vtX(76.5)), Y(598), L(vtX(104.5) - vtX(76.5)), L(18), "Sâu trong kho", 7, RED, True, box="FFFFFF")
FAST = [(5, 96), (7, 92), (9, 99), (11, 94), (13, 97), (15, 92)]
g2 = group(s, "Bước 2 – Hàng xuất nhiều nằm sâu")
for l, v in FAST:
    rect(g2.shapes, X(vtX(v)) - L(10), Y(laneY(l)) - L(10), L(20), L(20), RED, "FFFFFF", 1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.2)
callout(g2.shapes, vtX(55), laneY(17.0), 340, 36, "Hàng xuất nhiều nằm sâu", "xếp theo nhóm sản phẩm, gần khu nhập")
g3 = group(s, "Bước 3 – Quãng đường mỗi lần xuất")
for l, v in FAST:
    poly(g3.shapes, [(vtX(v) - 12, laneY(l)), (92, laneY(l))], RED, 1.5, dash=MSO_LINE.DASH)
g4 = bottom_banner(s, "Mỗi lần xuất, xe phải chạy gần hết chiều dài dãy", "Vị trí lưu kho chưa xét tần suất xuất hàng")
set_timing(s, [[(g1, 0)], [(g2, 0)], [(g3, 0)], [(g4, 0)]])

prs.save(DST)
print("saved", DST)
